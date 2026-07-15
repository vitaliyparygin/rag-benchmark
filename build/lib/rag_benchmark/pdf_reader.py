"""Document reading: abstract reader interface plus concrete implementations.

Despite the module name (kept to match the requested package layout), this
module hosts the DocumentReader interface and every built-in reader, not
just the PDF one. PDF was the first/primary reader so the module kept its
name; readers for other formats are simple enough to colocate here without
violating single-responsibility (each reader class still does exactly one
thing: turn one file format into a Document).
"""

from __future__ import annotations

from abc import ABC, abstractmethod
from pathlib import Path

from rag_benchmark.models import Document, DocumentFormat
from rag_benchmark.logging import get_logger
from rag_benchmark.documents.ids import stable_document_id

logger = get_logger("reader")


class DocumentReader(ABC):
    """Abstract interface for turning a file on disk into a Document."""

    @abstractmethod
    def read(self, path: Path) -> Document:
        """Read a file and return a normalized Document.

        Args:
            path: Path to the file to read.

        Returns:
            A populated Document with extracted text.

        Raises:
            ValueError: If the file cannot be parsed.
        """
        raise NotImplementedError


class PDFReader(DocumentReader):
    """Reads PDF files using pdfplumber."""

    def read(self, path: Path) -> Document:
        try:
            import pdfplumber
        except ImportError as exc:  # pragma: no cover - dependency guard
            raise RuntimeError(
                "pdfplumber is required to read PDF files. Install with "
                "`pip install pdfplumber`."
            ) from exc

        text_parts: list[str] = []
        page_count = 0
        try:
            with pdfplumber.open(path) as pdf:
                page_count = len(pdf.pages)
                for page in pdf.pages:
                    extracted = page.extract_text() or ""
                    text_parts.append(extracted)
        except Exception as exc:  # noqa: BLE001
            raise ValueError(f"Failed to read PDF {path}: {exc}") from exc

        text = "\n".join(text_parts)
        logger.info(
            "reader: %s : %d chars",
            path.name,
            len(text),
        )
        return Document(
            id=stable_document_id(path),
            path=path,
            filename=path.name,
            format=DocumentFormat.PDF,
            text=text,
            page_count=page_count,
            char_count=len(text),
        )


class DocxReader(DocumentReader):
    """Reads DOCX files using python-docx."""

    def read(self, path: Path) -> Document:
        try:
            import docx
        except ImportError as exc:  # pragma: no cover - dependency guard
            raise RuntimeError(
                "python-docx is required to read DOCX files. Install with "
                "`pip install python-docx`."
            ) from exc

        try:
            document = docx.Document(str(path))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        except Exception as exc:  # noqa: BLE001
            raise ValueError(f"Failed to read DOCX {path}: {exc}") from exc

        return Document(
            id=stable_document_id(path),
            path=path,
            filename=path.name,
            format=DocumentFormat.DOCX,
            text=text,
            char_count=len(text),
        )


class TxtReader(DocumentReader):
    """Reads plain text files."""

    def read(self, path: Path) -> Document:
        try:
            text = path.read_text(encoding="utf-8", errors="replace")
        except OSError as exc:
            raise ValueError(f"Failed to read TXT {path}: {exc}") from exc

        return Document(
            id=stable_document_id(path),
            path=path,
            filename=path.name,
            format=DocumentFormat.TXT,
            text=text,
            char_count=len(text),
        )


class MarkdownReader(DocumentReader):
    """Reads Markdown files as raw text (no HTML rendering needed for NLP use)."""

    def read(self, path: Path) -> Document:
        try:
            text = path.read_text(encoding="utf-8", errors="replace")
        except OSError as exc:
            raise ValueError(f"Failed to read Markdown {path}: {exc}") from exc

        return Document(
            id=stable_document_id(path),
            path=path,
            filename=path.name,
            format=DocumentFormat.MARKDOWN,
            text=text,
            char_count=len(text),
        )


class ReaderRegistry:
    """Maps DocumentFormat values to DocumentReader instances.

    New formats can be supported by calling `register` without touching
    existing reader code, keeping the system open for extension.
    """

    def __init__(self) -> None:
        self._readers: dict[DocumentFormat, DocumentReader] = {
            DocumentFormat.PDF: PDFReader(),
            DocumentFormat.DOCX: DocxReader(),
            DocumentFormat.TXT: TxtReader(),
            DocumentFormat.MARKDOWN: MarkdownReader(),
        }

    def register(self, fmt: DocumentFormat, reader: DocumentReader) -> None:
        """Register or override the reader used for a given format."""
        self._readers[fmt] = reader

    def get(self, fmt: DocumentFormat) -> DocumentReader:
        """Return the reader registered for a format.

        Raises:
            KeyError: If no reader is registered for the format.
        """
        if fmt not in self._readers:
            raise KeyError(f"No reader registered for format: {fmt}")
        return self._readers[fmt]

    def read(self, path: Path, fmt: DocumentFormat) -> Document:
        """Convenience helper: look up the right reader and read the file."""
        reader = self.get(fmt)
        logger.debug("Reading %s with %s", path, type(reader).__name__)
        return reader.read(path)
